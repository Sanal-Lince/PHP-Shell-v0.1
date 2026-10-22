"""DOCX and PDF export helpers for PHP2Scope scan results."""
from __future__ import annotations
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape
from typing import Any

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether


def _safe(value: Any, default: str = "-", maxlen: int = 500) -> str:
    if value is None or value == "": return default
    s = str(value).replace("\x00", " ")
    return s[:maxlen]


def _assessment(d: dict[str, Any]) -> dict[str, Any]:
    return d.get("version_assessment") or {}


def export_docx(data: dict[str, Any], output: str | Path) -> Path:
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65); sec.left_margin = Inches(.7); sec.right_margin = Inches(.7)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"; normal.font.size = Pt(9)
    normal.font.color.rgb = RGBColor(42, 52, 65)
    if "List Bullet" in doc.styles:
        doc.styles["List Bullet"].paragraph_format.space_before = Pt(0)
        doc.styles["List Bullet"].paragraph_format.space_after = Pt(1)
        doc.styles["List Bullet"].paragraph_format.line_spacing = 1.0
    for name, size, color in [("Title", 24, (17, 55, 90)), ("Heading 1", 15, (17, 75, 110)), ("Heading 2", 11, (34, 81, 116))]:
        st = doc.styles[name]; st.font.name = "Aptos Display"; st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = RGBColor(*color)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("PHP2Scope")
    r.bold = True; r.font.size = Pt(25); r.font.color.rgb = RGBColor(17, 75, 110)
    p = doc.add_paragraph("Passive PHP CVE Assessment Report")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    p = doc.add_paragraph("CVE-2026-109618 · CVE-2026-109619")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.size = Pt(10); p.runs[0].font.color.rgb = RGBColor(85, 105, 125)

    doc.add_heading("1. Assessment summary", level=1)
    summary = data.get("summary", {})
    va = _assessment(data)
    rows = [
        ("Target", _safe(data.get("target_url"))),
        ("Scan ID", _safe(data.get("scan_id"))),
        ("Generated (UTC)", _safe(data.get("created_utc"))),
        ("Access mode", _safe(data.get("access_mode"))),
        ("Validation depth", _safe(data.get("validation_depth"))),
        ("PHP version result", _safe(va.get("label", "PHP version not determined"))),
        ("Observed / supplied version", _safe(va.get("version"))),
        ("Assessment confidence", _safe(va.get("confidence"))),
        ("HTTP requests", _safe(summary.get("http_requests", 0))),
        ("Candidate endpoints", _safe(summary.get("endpoint_count", 0))),
        ("Candidate parameters", _safe(summary.get("parameter_count", 0))),
    ]
    _docx_table(doc, ["Field", "Value"], rows)
    doc.add_paragraph(_safe(va.get("detail", "No exact PHP version could be established from current evidence."), maxlen=1200))

    doc.add_heading("2. CVE context", level=1)
    doc.add_paragraph("CVE-2026-109618 concerns out-of-memory handling where request data may be partially initialized before cleanup. CVE-2026-109619 concerns stale request-owned internal state after its backing memory is freed. The PHP security team published fixes in PHP 8.5.12, 8.4.27, 8.3.36 and 8.2.35. Vendor packages can backport fixes, so package provenance matters.")
    doc.add_paragraph("For CVE-2026-109618 and CVE-2026-109619, the result above is version triage only and does not confirm remote exploitability. A missing PHP banner is not proof of safety.")

    doc.add_heading("3. Technology fingerprints", level=1)
    tech = data.get("technology_findings", [])
    if tech:
        _docx_table(doc, ["Technology", "Confidence", "Evidence"], [(_safe(x.get("technology")), _safe(x.get("confidence")), _safe(x.get("evidence"), maxlen=300)) for x in tech])
    else: doc.add_paragraph("No technology fingerprints were identified by enabled modules.")

    doc.add_heading("4. HTTP observations", level=1)
    obs = data.get("http_observations", [])
    if obs:
        _docx_table(doc, ["Method", "URL", "Status", "Time (ms)", "Notes"], [(_safe(x.get("method")), _safe(x.get("url"), maxlen=100), _safe(x.get("status_code")), _safe(x.get("elapsed_ms")), _safe(x.get("error") or x.get("redirect_location") or "")) for x in obs])
        header_map = obs[0].get("headers") or {}
        display_headers = {k: v for k, v in header_map.items() if k.lower() in {"server", "x-powered-by", "content-type", "content-security-policy", "strict-transport-security", "x-content-type-options", "x-frame-options", "referrer-policy", "permissions-policy", "cache-control", "access-control-allow-origin", "access-control-allow-credentials", "via", "x-cache"}}
        hp = doc.add_paragraph()
        hr = hp.add_run("Baseline response headers (allowlisted): ")
        hr.bold = True
        hp.add_run("; ".join(f"{k}={v}" for k, v in display_headers.items()) if display_headers else "none observed")
        hp.paragraph_format.space_after = Pt(2)
    else: doc.add_paragraph("No HTTP observations are available.")
    notes = data.get("evidence_notes", [])
    report_notes = [n for n in notes if not n.startswith(("HTTP response header contains a PHP-related token", "PHPSESSID cookie name observed", "HTTP finding:"))]
    if report_notes:
        doc.add_heading("Evidence notes", level=2)
        for note in report_notes[:100]: doc.add_paragraph(_safe(note, maxlen=500), style="List Bullet")

    doc.add_heading("5. Candidate endpoint inventory", level=1)
    endpoints = data.get("endpoint_inventory", [])
    if endpoints:
        _docx_table(doc, ["Method", "Path", "Query names", "Discovered from"], [(_safe(x.get("method")), _safe(x.get("path"), maxlen=140), _safe(", ".join(x.get("query_parameter_names", [])), maxlen=200), _safe(x.get("source"), maxlen=100)) for x in endpoints[:250]])
    else: doc.add_paragraph("No endpoint inventory items were discovered or the module was not selected.")

    doc.add_heading("6. Candidate parameter inventory", level=1)
    params = data.get("parameter_inventory", [])
    if params:
        _docx_table(doc, ["Path", "Parameter", "Location", "Method", "Tested"], [(_safe(x.get("path"), maxlen=120), _safe(x.get("parameter")), _safe(x.get("location")), _safe(x.get("method", "GET")), "No") for x in params[:250]])
    else: doc.add_paragraph("No candidate parameters were identified. PHP2Scope does not submit forms or test parameters in v1.")

    doc.add_heading("7. Limitations and recommended next actions", level=1)
    for line in [
        "Treat the result as triage evidence, not proof of exploitability or proof of absence.",
        "Confirm the runtime used by the internet-facing virtual host or worker; CLI PHP and web-serving PHP may differ.",
        "For versions below the upstream fixed release, verify OS/vendor package changelogs and backport status.",
        "Apply a vendor-supported patched PHP release promptly and rebuild/redeploy containers or serverless images where applicable.",
        "Deeper validation should be performed only in an isolated, authorized test environment. This v1 sends only bounded GET/HEAD requests and never submits forms or payloads.",
    ]: doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("8. Sources", level=1)
    source_labels = [
        ("PHP Security Advisory (php.net; 21 Oct 2026)", "https://www.php.net/archive/2026.php"),
        ("Google Cloud analysis (21 Oct 2026)", "https://cloud.google.com/blog/products/identity-security/responding-to-core-php-vulnerabilities-cve-2026-109618-cve-2026-109619"),
        ("Canadian Centre for Cyber Security alert (21 Oct 2026)", "https://www.cyber.gc.ca/en/alerts-advisories/al26-026-critical-vulnerabilities-impacting-php-cve-2026-109618-cve-2026-109619"),
    ]
    source_p = doc.add_paragraph()
    source_p.paragraph_format.space_before = Pt(0)
    source_p.paragraph_format.space_after = Pt(0)
    source_p.paragraph_format.line_spacing = 1.0
    for i, (label, url) in enumerate(source_labels):
        if i:
            source_p.add_run(" | ")
        _add_hyperlink(source_p, label, url)
    for run in source_p.runs:
        run.font.size = Pt(8.5)
    footer = sec.footer.paragraphs[0]
    footer.text = "PHP2Scope v0.1 | Passive-only assessment | Handle output as security-sensitive evidence"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        run.font.size = Pt(8); run.font.color.rgb = RGBColor(100, 110, 120)
    doc.save(output)
    return output


def _add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "0563C1"); rPr.append(color)
    underline = OxmlElement("w:u"); underline.set(qn("w:val"), "single"); rPr.append(underline)
    run.append(rPr)
    text_el = OxmlElement("w:t"); text_el.text = text; run.append(text_el)
    hyperlink.append(run); paragraph._p.append(hyperlink)


def _docx_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Shading Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]; cell.text = str(h); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in cell.paragraphs[0].runs: r.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = _safe(val, maxlen=500)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    doc.add_paragraph("")


def export_pdf(data: dict[str, Any], output: str | Path) -> Path:
    output = Path(output); output.parent.mkdir(parents=True, exist_ok=True)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ReportTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=22, leading=26, textColor=colors.HexColor("#114B6E"), alignment=TA_CENTER, spaceAfter=4))
    styles.add(ParagraphStyle(name="SubtitleCenter", parent=styles["Normal"], fontName="Helvetica", fontSize=10, textColor=colors.HexColor("#526779"), alignment=TA_CENTER, spaceAfter=12))
    styles.add(ParagraphStyle(name="SectionBlue", parent=styles["Heading1"], fontSize=14, leading=18, textColor=colors.HexColor("#114B6E"), spaceBefore=12, spaceAfter=6))
    styles.add(ParagraphStyle(name="SmallCell", parent=styles["BodyText"], fontSize=7.5, leading=9, wordWrap="CJK"))
    styles.add(ParagraphStyle(name="BodySmall", parent=styles["BodyText"], fontSize=8.5, leading=12))
    story = [Paragraph("PHP2Scope", styles["ReportTitle"]), Paragraph("Passive PHP CVE Assessment Report", styles["SubtitleCenter"]), Paragraph("CVE-2026-109618 · CVE-2026-109619", styles["SubtitleCenter"])]
    va = _assessment(data); sm = data.get("summary", {})
    story.append(Paragraph("1. Assessment summary", styles["SectionBlue"]))
    summary_rows = [
        ["Field", "Value"], ["Target", _safe(data.get("target_url"))], ["Scan ID", _safe(data.get("scan_id"))],
        ["Generated (UTC)", _safe(data.get("created_utc"))], ["Access mode", _safe(data.get("access_mode"))],
        ["Validation", _safe(data.get("validation_depth"))], ["PHP version", _safe(va.get("version"))],
        ["Assessment", _safe(va.get("label", "Not determined"))], ["Confidence", _safe(va.get("confidence"))],
        ["HTTP requests", _safe(sm.get("http_requests", 0))], ["Endpoints / parameters", f"{sm.get('endpoint_count', 0)} / {sm.get('parameter_count', 0)}"],
    ]
    story.append(_pdf_table(summary_rows, [42*mm, 125*mm]))
    story.append(Spacer(1, 5*mm))
    story.append(Paragraph(escape(_safe(va.get("detail", "No exact PHP version could be established."), maxlen=1200)), styles["BodySmall"]))
    story.append(Paragraph("2. CVE context", styles["SectionBlue"]))
    story.append(Paragraph("CVE-2026-109618 concerns out-of-memory handling with partially initialized request data. CVE-2026-109619 concerns stale request-owned internal state after backing memory is freed. Upstream fixed releases are PHP 8.5.12, 8.4.27, 8.3.36 and 8.2.35. Vendor packages may backport fixes.", styles["BodySmall"]))
    story.append(Spacer(1, 2*mm))
    story.append(Paragraph("For CVE-2026-109618 and CVE-2026-109619, this is version triage only and does not confirm remote exploitability. A missing PHP banner is not proof of safety.", styles["BodySmall"]))
    story.append(Paragraph("3. Technology fingerprints", styles["SectionBlue"]))
    tech = data.get("technology_findings", [])
    story.append(_pdf_table([["Technology", "Confidence", "Evidence"]] + [[_safe(x.get("technology")), _safe(x.get("confidence")), _safe(x.get("evidence"), maxlen=220)] for x in tech] if tech else [["Technology", "Confidence", "Evidence"], ["None", "-", "No fingerprints found or module not selected."]], [35*mm, 27*mm, 105*mm]))
    story.append(Paragraph("4. HTTP observations", styles["SectionBlue"]))
    obs = data.get("http_observations", [])
    story.append(_pdf_table([["Method", "URL", "Status", "ms", "Error/redirect"]] + [[_safe(x.get("method")), _safe(x.get("url"), maxlen=85), _safe(x.get("status_code")), _safe(x.get("elapsed_ms")), _safe(x.get("error") or x.get("redirect_location") or "", maxlen=100)] for x in obs] if obs else [["Method", "URL", "Status", "ms", "Error/redirect"]], [16*mm, 65*mm, 16*mm, 13*mm, 57*mm]))
    if obs:
        header_map = obs[0].get("headers") or {}
        display_headers = {k: v for k, v in header_map.items() if k.lower() in {"server", "x-powered-by", "content-type", "content-security-policy", "strict-transport-security", "x-content-type-options", "x-frame-options", "referrer-policy", "permissions-policy", "cache-control", "access-control-allow-origin", "access-control-allow-credentials", "via", "x-cache"}}
        header_block = [Paragraph("Baseline response headers (allowlisted)", styles["Heading2"]), _pdf_table([["Header", "Observed value"]] + [[k, v] for k, v in display_headers.items()] if display_headers else [["Header", "Observed value"], ["-", "No allowlisted headers observed."]], [45*mm, 122*mm])]
        story.append(KeepTogether(header_block))
    notes = data.get("evidence_notes", [])
    report_notes = [n for n in notes if not n.startswith(("HTTP response header contains a PHP-related token", "PHPSESSID cookie name observed", "HTTP finding:"))]
    if report_notes:
        story.append(Paragraph("Evidence notes", styles["Heading2"]))
        for n in report_notes[:60]: story.append(Paragraph("• " + escape(_safe(n, maxlen=400)), styles["BodySmall"]))
    story.append(Paragraph("5. Candidate endpoint inventory", styles["SectionBlue"]))
    eps = data.get("endpoint_inventory", [])
    endpoint_rows = [["Method", "Path", "Query names", "Source"]] + [[_safe(x.get("method")), _safe(x.get("path"), maxlen=110), _safe(", ".join(x.get("query_parameter_names", [])), maxlen=80), _safe(x.get("source"), maxlen=70)] for x in eps[:250]]
    story.append(_pdf_table(endpoint_rows, [18*mm, 65*mm, 42*mm, 42*mm]))
    story.append(Paragraph("6. Candidate parameter inventory", styles["SectionBlue"]))
    params = data.get("parameter_inventory", [])
    param_rows = [["Path", "Parameter", "Location", "Method", "Tested"]] + [[_safe(x.get("path"), maxlen=75), _safe(x.get("parameter")), _safe(x.get("location")), _safe(x.get("method", "GET")), "No"] for x in params[:250]]
    story.append(_pdf_table(param_rows, [55*mm, 38*mm, 35*mm, 18*mm, 21*mm]))
    story.append(Paragraph("7. Limitations and next actions", styles["SectionBlue"]))
    for line in ["Treat results as triage evidence, not proof of exploitability or proof of absence.", "Confirm the version used by the internet-facing web-serving runtime; CLI PHP can differ.", "Verify vendor backports for old branches and package builds.", "Apply supported vendor patches and rebuild/redeploy image-based deployments.", "This v1 uses bounded GET/HEAD requests only: no form submission, fuzzing, parameters tested, or exploit attempts."]:
        story.append(Paragraph("• " + escape(line), styles["BodySmall"]))
    story.append(Paragraph("8. Sources", styles["SectionBlue"]))
    for u in data.get("advisory_sources", []): story.append(Paragraph(escape(_safe(u, maxlen=500)), styles["BodySmall"]))
    def footer(canvas, doc):
        canvas.saveState(); canvas.setFont("Helvetica", 7); canvas.setFillColor(colors.HexColor("#64748b"))
        canvas.drawString(16*mm, 9*mm, "PHP2Scope v0.1 | Passive-only assessment | Security-sensitive evidence")
        canvas.drawRightString(A4[0]-16*mm, 9*mm, f"Page {doc.page}")
        canvas.restoreState()
    pdf = SimpleDocTemplate(str(output), pagesize=A4, rightMargin=16*mm, leftMargin=16*mm, topMargin=15*mm, bottomMargin=16*mm, title="PHP2Scope Assessment Report", author="PHP2Scope")
    pdf.build(story, onFirstPage=footer, onLaterPages=footer)
    return output


def _pdf_table(rows, widths):
    cooked = []
    header_style = ParagraphStyle(name="TableHeaderWhite", fontName="Helvetica-Bold", fontSize=7.4, leading=8.8, textColor=colors.white)
    for ri, row in enumerate(rows):
        cell_style = header_style if ri == 0 else ParagraphStyle(name=f"Cell{ri}_{len(cooked)}", fontName="Helvetica", fontSize=7.2, leading=8.5, textColor=colors.HexColor("#263747"))
        cooked.append([Paragraph(escape(_safe(cell, maxlen=600)), cell_style) for cell in row])
    table = Table(cooked, colWidths=widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#114B6E")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("GRID", (0,0), (-1,-1), 0.35, colors.HexColor("#d3dde5")),
        ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#f3f7fa")]),
        ("LEFTPADDING", (0,0), (-1,-1), 5), ("RIGHTPADDING", (0,0), (-1,-1), 5),
        ("TOPPADDING", (0,0), (-1,-1), 4), ("BOTTOMPADDING", (0,0), (-1,-1), 4),
    ]))
    return table
